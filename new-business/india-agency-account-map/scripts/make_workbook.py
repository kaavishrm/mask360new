"""Build the India agency account map and poaching target list workbook."""
import re, sys, os
from collections import Counter, defaultdict
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from load import load, dedupe, CONF_RANK
from canon import canon_agency
from brand_merge import brand_key
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "India_Agency_Account_Map_Sep2026.xlsx")
ASOF = "29 Sep 2026"

# ---------- scoring rules (documented on the Read me sheet) ----------
FIT = {
    "Luxury hospitality & travel": 4, "Jewellery & watches": 4, "Fashion & luxury retail": 4,
    "Real estate": 4, "Beauty & personal care": 3, "Alcobev": 3, "Automotive": 3,
    "Wellness, fitness & sport": 3, "BFSI & wealth": 2, "Premium F&B & QSR": 2, "Aviation": 2,
    "Consumer tech & D2C": 1, "E-commerce & quick commerce": 1, "FMCG": 1, "Telecom & media": 1,
    "Healthcare": 1, "Other": 1, "Education": 0,
}
TIER_PTS = {"Luxury": 3, "Premium": 2, "Mass": 0}
TIER_RANK = {"Luxury": 3, "Premium": 2, "Mass": 1}

CLIENT_BRANDS = r"anantara|\bdior\b|ferrari|zorae|zor.e|kalyan jewel|\ba29\b|atenx|burjeel"
CLIENT_PARENTS = r"l'or[eé]al|loreal"

UNVERIFIED = [r"bluestone.*ipo|ipo.*bluestone", r"hul.*minimalist|minimalist.*hul", r"bira.*(money|debt|cash|financ)",
              r"good glamm", r"kenvue", r"smbc", r"opella", r"kwality wall"]

PAT = {
    "Pitch or review live": r"call(ed|s)? (a |for )?(creative |digital |social |pr )?pitch|pitch (called|invited|underway|is live|on)|in review|under review|account review|review of (its|the)|invit(ed|es|ing) agencies|\brfp\b|open (public )?pitch|creative review|put .{0,30} review",
    "Leadership change": r"new (cmo|chief marketing|head of marketing|marketing head|marketing director|marketing lead|brand head|brand director|ceo|md|country manager|chief|president|leadership)|\bcmo\b.{0,40}(left|exit|quit|join|appoint)|appointed|joins as|joined as|elevated|stepped (down|back)|\bleft\b|\bexit|departed|resign|moved to",
    "Incumbent turmoil": r"retired|absorbed|merg|folded|acqui|for sale|explor\w* (a )?sale|insolven|layoff|job cuts|cut \d+|tax (probe|search)|restructur|dissolved|wound down|lapsed|in limbo|no confirmed|churn|lawsuit|suing|disabled",
    "Split roster": r"\bsplit\b|across (two|three|four|several) agencies|multiple agencies|several agencies",
    "Long tenure": r"since (19\d\d|200\d|201[0-8])|\b(\d{2}|[5-9])\+? years|decade|fatigue|\b(5th|fifth|sixth|seventh) (year|term)",
    "Fresh budget or launch": r"\blaunch|\bipo\b|rais(ed|ing|e)|funding|expan|new hotel|new project|sponsor|relaunch|new brand|entry into|debut",
    "Renewal due": r"one[- ]year|renewal|up for renewal",
    "Open mandates": r"unassigned|no agency (is )?(named|found|credited|visible|appointed)|(are|is|look|looks|remain|remains|left) open|openings?\b|single[- ]channel|without a confirmed|no confirmed (new )?home",
}

def triggers_for(text):
    t = text.lower()
    return [k for k, p in PAT.items() if re.search(p, t)]

def since_year(s):
    m = re.search(r"(19|20)\d\d", s or "")
    return int(m.group(0)) if m else None

# ---------- load ----------
raw, bad = load()
if bad:
    print("bad lines:", bad)
recs = dedupe(raw)
for r in recs:
    r["agency_raw"] = r["agency"]
    r["agency"], r["agency_group"] = canon_agency(r["agency"], r["agency_group"])
    if r["agency"] == "Unknown or in-house":
        r["agency_type"] = "Unknown"
    k, disp = brand_key(r["brand"])
    r["_bkey"], r["_bdisp"] = k, disp
    sig = r["poach_signals"]
    low = (r["brand"] + " " + sig).lower()
    if sig and any(re.search(p, low) for p in UNVERIFIED):
        r["poach_signals"] = sig + " [unverified, check before outreach]"

# dedupe again after canonicalisation (same brand, agency, mandate)
seen = {}
for r in recs:
    key = (r["_bkey"], r["agency"], r["mandate"])
    if key in seen:
        cur = seen[key]
        for u in r.get("_sources", []):
            if u and u not in cur["_sources"]:
                cur["_sources"].append(u)
        if (CONF_RANK.get(r["confidence"], 0), r["source_date"]) > (CONF_RANK.get(cur["confidence"], 0), cur["source_date"]):
            r["_sources"] = cur["_sources"]
            if cur["poach_signals"] and cur["poach_signals"] not in r["poach_signals"]:
                r["poach_signals"] = (r["poach_signals"] + "; " + cur["poach_signals"]).strip("; ")
            seen[key] = r
        elif r["poach_signals"] and r["poach_signals"] not in cur["poach_signals"]:
            cur["poach_signals"] = (cur["poach_signals"] + "; " + r["poach_signals"]).strip("; ")
    else:
        seen[key] = r
recs = list(seen.values())
print("records after canonical dedupe:", len(recs))

# brand display names
groups = defaultdict(list)
for r in recs:
    groups[r["_bkey"]].append(r)
for k, rs in groups.items():
    disp = next((r["_bdisp"] for r in rs if r["_bdisp"]), None)
    if not disp:
        c = Counter(r["brand"] for r in rs)
        disp = sorted(c.items(), key=lambda x: (-x[1], len(x[0])))[0][0]
    for r in rs:
        r["brand_display"] = disp

# ---------- brand level target rows ----------
STATUS_ORDER = {"In review": 0, "Won recently": 1, "Active": 2, "Unclear": 3, "Recently lost": 4}
targets = []
for k, rs in groups.items():
    rs.sort(key=lambda r: (STATUS_ORDER.get(r["status"], 9), -CONF_RANK.get(r["confidence"], 0), r["source_date"]), reverse=False)
    disp = rs[0]["brand_display"]
    parent = Counter(r["parent_company"] for r in rs if r["parent_company"]).most_common(1)
    parent = parent[0][0] if parent else ""
    cat = Counter(r["category"] for r in rs).most_common(1)[0][0]
    tier = max((r["brand_tier"] for r in rs), key=lambda t: TIER_RANK.get(t, 0))
    known = [r for r in rs if r["agency"] != "Unknown or in-house"]
    current = [r for r in known if r["status"] != "Recently lost"]
    lost = [r for r in known if r["status"] == "Recently lost"]
    def fmt(r):
        bits = [r["mandate"]]
        y = since_year(r["since"])
        if y:
            bits.append(f"since {y}")
        s = f"{r['agency']} ({', '.join(bits)})"
        if r["status"] == "Unclear":
            s += " unconfirmed"
        if r["status"] == "In review":
            s += " IN REVIEW"
        return s
    cur_txt = "; ".join(dict.fromkeys(fmt(r) for r in current)) or "None credited (in-house or unknown)"
    lost_txt = "; ".join(dict.fromkeys(f"{r['agency']} ({r['mandate']})" for r in lost))
    sigs = list(dict.fromkeys(s.strip() for r in rs for s in re.split(r";\s*", r["poach_signals"]) if s.strip()))
    sig_txt = "; ".join(sigs)
    trig = triggers_for(sig_txt)
    if any(r["status"] == "In review" for r in rs) and "Pitch or review live" not in trig:
        trig.insert(0, "Pitch or review live")
    if not current:
        trig.append("No agency credited")
    if len({r["agency"] for r in current}) >= 3 and "Split roster" not in trig:
        trig.append("Split roster")
    ys = [since_year(r["since"]) for r in current if since_year(r["since"])]
    if ys and min(ys) <= 2018 and "Long tenure" not in trig:
        trig.append("Long tenure")
    newest = max(rs, key=lambda r: r["source_date"])
    ts = set(trig)
    timing = 0
    timing += 3 if "Pitch or review live" in ts else 0
    timing += 2 if "Leadership change" in ts else 0
    timing += 2 if "Incumbent turmoil" in ts else 0
    timing += 2 if "Open mandates" in ts else 0
    timing += 1 if "No agency credited" in ts else 0
    timing += 1 if {"Split roster", "Long tenure", "Fresh budget or launch", "Renewal due"} & ts else 0
    timing = min(3, timing)
    honeymoon = any(r["status"] == "Won recently" and r["source_date"] >= "2026-03" for r in current) \
        and not {"Pitch or review live", "Open mandates"} & ts
    if honeymoon:
        timing = 0
        trig.append("Just moved (honeymoon)")
    blob = (disp + " " + " ".join(r["brand"] for r in rs)).lower()
    client = ""
    if re.search(CLIENT_BRANDS, blob):
        client = "Existing Mask360 client"
    elif re.search(CLIENT_PARENTS, (parent + " " + blob).lower()):
        client = "Mask360 client group (L'Oreal): expand"
    srcs = list(dict.fromkeys(u for r in sorted(rs, key=lambda r: (-CONF_RANK.get(r["confidence"], 0), r["source_date"])) for u in r.get("_sources", []) if u))
    best = max(rs, key=lambda r: (CONF_RANK.get(r["confidence"], 0), r["source_date"]))
    groups_held = sorted({r["agency_group"] for r in current})
    targets.append({
        "brand": disp, "parent": parent, "category": cat, "tier": tier,
        "fit": FIT.get(cat, 1), "tierpts": TIER_PTS.get(tier, 0), "timing": timing,
        "current": cur_txt, "groups": ", ".join(groups_held) if groups_held else "Unknown",
        "lost": lost_txt, "triggers": ", ".join(dict.fromkeys(trig)), "signals": sig_txt,
        "evidence": best["evidence"], "latest": newest["source_date"], "confidence": best["confidence"],
        "sources": srcs[:3], "client": client, "n": len(rs),
    })

for t in targets:
    t["score"] = t["fit"] + t["tierpts"] + t["timing"]
targets.sort(key=lambda t: (t["client"] != "", -t["score"], -TIER_RANK.get(t["tier"], 0), -t["timing"], t["brand"].lower()))
print("brands:", len(targets))
print("A:", sum(1 for t in targets if not t["client"] and t["score"] >= 9),
      "B:", sum(1 for t in targets if not t["client"] and 7 <= t["score"] < 9))

# ---------- workbook ----------
FONT = "Arial"
HDR_FILL = PatternFill("solid", start_color="1F1F1F")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
BLUE = Font(name=FONT, size=10, color="0000FF")
LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
YELLOW = PatternFill("solid", start_color="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(bottom=thin)

wb = Workbook()

def header(ws, cols, row=1):
    for i, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font, c.fill = HDR_FONT, HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=3)

def put(ws, r, c, v, font=BODY, align=TOP, fill=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font, cell.alignment = font, align
    if fill:
        cell.fill = fill
    return cell

def link(ws, r, c, url):
    if not url:
        return put(ws, r, c, "")
    cell = put(ws, r, c, url, LINK)
    cell.hyperlink = url
    return cell

# ----- All relationships (built first so other sheets can reference it) -----
rel = wb.active
rel.title = "All relationships"
REL_COLS = [("Agency (current name)", 30), ("Agency group", 13), ("Agency type", 20), ("Brand", 30),
            ("Brand as reported", 30), ("Parent company", 24), ("Category", 22), ("Brand tier", 10),
            ("Mandate", 12), ("Since", 9), ("Status", 13), ("Evidence", 60), ("Poaching signals", 50),
            ("Source", 45), ("Source date", 10), ("Confidence", 11), ("Agency as reported", 28), ("Research stream", 14)]
header(rel, REL_COLS)
GROUP_ORDER = ["WPP", "Publicis", "Omnicom", "Dentsu", "Havas", "Cheil", "Hakuhodo", "Independent", "Other", "Unknown"]
recs.sort(key=lambda r: (GROUP_ORDER.index(r["agency_group"]) if r["agency_group"] in GROUP_ORDER else 8,
                         r["agency"].lower(), -TIER_RANK.get(r["brand_tier"], 0), r["brand_display"].lower()))
STREAM = {"wpp": "WPP", "publicis": "Publicis", "omnicom_a": "Omnicom 1", "omnicom_b": "Omnicom 2 (ex IPG)",
          "dentsu_havas": "Dentsu, Havas, Asian", "indie_creative": "Indie creative", "indie_digital": "Indie digital",
          "luxury_pr_exp": "Luxury, PR, events", "in_play": "In play", "brands_a": "Brands: travel, realty, auto, BFSI",
          "brands_b": "Brands: jewellery, fashion, beauty, alcobev"}
for i, r in enumerate(recs, 2):
    vals = [r["agency"], r["agency_group"], r["agency_type"], r["brand_display"], r["brand"], r["parent_company"],
            r["category"], r["brand_tier"], r["mandate"], r["since"], r["status"], r["evidence"], r["poach_signals"]]
    for j, v in enumerate(vals, 1):
        put(rel, i, j, v, align=WRAP if j in (12, 13) else TOP)
    link(rel, i, 14, r["source_url"])
    put(rel, i, 15, r["source_date"])
    put(rel, i, 16, r["confidence"])
    put(rel, i, 17, r["agency_raw"])
    put(rel, i, 18, STREAM.get(r["_file"], r["_file"]))
REL_LAST = len(recs) + 1
rel.auto_filter.ref = f"A1:{get_column_letter(len(REL_COLS))}{REL_LAST}"

# ----- Target list -----
tl = wb.create_sheet("Target list", 0)
TL_COLS = [("Rank", 6), ("Brand", 28), ("Priority", 9), ("Score (of 10)", 8), ("Fit pts", 6), ("Tier pts", 6),
           ("Timing pts (edit)", 8), ("Category", 20), ("Brand tier", 9), ("Parent company", 22),
           ("Current agency and mandate", 48), ("Held by group", 14), ("Recently lost by", 26), ("Why now (triggers)", 30),
           ("Poaching signals", 60), ("Key evidence", 50), ("Latest source", 10), ("Confidence", 10),
           ("Source 1", 36), ("Source 2", 36), ("Mask360 relationship", 20),
           ("Owner", 12), ("Outreach status", 14), ("Next step", 24), ("Notes", 30)]
header(tl, TL_COLS)
for i, t in enumerate(targets, 2):
    put(tl, i, 1, i - 1)
    put(tl, i, 2, t["brand"], BOLD)
    put(tl, i, 3, f'=IF(U{i}<>"","Client",IF(D{i}>=9,"A",IF(D{i}>=7,"B","C")))')
    put(tl, i, 4, f"=E{i}+F{i}+G{i}")
    put(tl, i, 5, t["fit"])
    put(tl, i, 6, t["tierpts"])
    put(tl, i, 7, t["timing"], BLUE)
    put(tl, i, 8, t["category"])
    put(tl, i, 9, t["tier"])
    put(tl, i, 10, t["parent"], align=WRAP)
    put(tl, i, 11, t["current"], align=WRAP)
    put(tl, i, 12, t["groups"], align=WRAP)
    put(tl, i, 13, t["lost"], align=WRAP)
    put(tl, i, 14, t["triggers"], align=WRAP)
    put(tl, i, 15, t["signals"], align=WRAP)
    put(tl, i, 16, t["evidence"], align=WRAP)
    put(tl, i, 17, t["latest"])
    put(tl, i, 18, t["confidence"])
    link(tl, i, 19, t["sources"][0] if t["sources"] else "")
    link(tl, i, 20, t["sources"][1] if len(t["sources"]) > 1 else "")
    put(tl, i, 21, t["client"])
    for c in (22, 23, 24, 25):
        put(tl, i, c, "Not started" if c == 23 else None, fill=YELLOW)
TL_LAST = len(targets) + 1
tl.auto_filter.ref = f"A1:{get_column_letter(len(TL_COLS))}{TL_LAST}"
tl.freeze_panes = "C2"
dv = DataValidation(type="list", formula1='"Not started,Researching,Contacted,Meeting set,Pitching,Won,Parked"', allow_blank=True)
tl.add_data_validation(dv)
dv.add(f"W2:W{TL_LAST}")
dvt = DataValidation(type="whole", operator="between", formula1="0", formula2="3", allow_blank=False)
tl.add_data_validation(dvt)
dvt.add(f"G2:G{TL_LAST}")
tl.conditional_formatting.add(f"C2:C{TL_LAST}", CellIsRule(operator="equal", formula=['"A"'], fill=PatternFill("solid", start_color="C6EFCE"), font=Font(name=FONT, bold=True, color="006100")))
tl.conditional_formatting.add(f"C2:C{TL_LAST}", CellIsRule(operator="equal", formula=['"B"'], fill=PatternFill("solid", start_color="FFEB9C"), font=Font(name=FONT, bold=True, color="9C5700")))
tl.conditional_formatting.add(f"C2:C{TL_LAST}", CellIsRule(operator="equal", formula=['"Client"'], fill=PatternFill("solid", start_color="DDEBF7"), font=Font(name=FONT, color="1F4E78")))

# ----- In play now -----
ip = wb.create_sheet("In play now", 1)
IP_COLS = [("Date", 9), ("Brand", 28), ("Category", 20), ("Brand tier", 9), ("What is happening", 60),
           ("Agency involved", 28), ("Agency status", 13), ("Mandate", 11), ("Evidence", 55), ("Source", 40), ("Confidence", 10)]
header(ip, IP_COLS)
play = [r for r in recs if r["_file"] == "in_play" or r["status"] == "In review"]
play.sort(key=lambda r: (r["source_date"], TIER_RANK.get(r["brand_tier"], 0)), reverse=True)
for i, r in enumerate(play, 2):
    vals = [r["source_date"], r["brand_display"], r["category"], r["brand_tier"], r["poach_signals"], r["agency"],
            r["status"], r["mandate"], r["evidence"]]
    for j, v in enumerate(vals, 1):
        put(ip, i, j, v, BOLD if j == 2 else BODY, WRAP if j in (5, 9) else TOP)
    link(ip, i, 10, r["source_url"])
    put(ip, i, 11, r["confidence"])
ip.auto_filter.ref = f"A1:K{len(play) + 1}"

# ----- Agency roster -----
ar = wb.create_sheet("Agency roster", 2)
AR_COLS = [("Agency (current name)", 32), ("Group", 12), ("Type", 20), ("Relationships found", 11),
           ("Luxury accounts", 9), ("Premium accounts", 9), ("Luxury and premium accounts held", 80),
           ("Other accounts held", 60), ("Recently lost", 34), ("Also reported as", 34)]
header(ar, AR_COLS)
by_ag = defaultdict(list)
for r in recs:
    if r["agency"] != "Unknown or in-house":
        by_ag[r["agency"]].append(r)
ag_rows = sorted(by_ag.items(), key=lambda kv: (GROUP_ORDER.index(kv[1][0]["agency_group"]) if kv[1][0]["agency_group"] in GROUP_ORDER else 8, -len(kv[1]), kv[0].lower()))
for i, (ag, rs) in enumerate(ag_rows, 2):
    cur = [r for r in rs if r["status"] != "Recently lost"]
    lux = sorted({r["brand_display"] for r in cur if r["brand_tier"] in ("Luxury", "Premium")}, key=lambda b: (-max(TIER_RANK[r["brand_tier"]] for r in cur if r["brand_display"] == b), b.lower()))
    lux_txt = ", ".join(b + (" (L)" if any(r["brand_tier"] == "Luxury" for r in cur if r["brand_display"] == b) else "") for b in lux)
    other = sorted({r["brand_display"] for r in cur if r["brand_tier"] not in ("Luxury", "Premium")}, key=str.lower)
    lost = sorted({r["brand_display"] for r in rs if r["status"] == "Recently lost"}, key=str.lower)
    aka = sorted({r["agency_raw"] for r in rs if r["agency_raw"] != ag}, key=str.lower)
    put(ar, i, 1, ag, BOLD)
    put(ar, i, 2, rs[0]["agency_group"])
    put(ar, i, 3, Counter(r["agency_type"] for r in rs).most_common(1)[0][0])
    put(ar, i, 4, f"=COUNTIF('All relationships'!$A$2:$A${REL_LAST},A{i})")
    put(ar, i, 5, f"=COUNTIFS('All relationships'!$A$2:$A${REL_LAST},A{i},'All relationships'!$H$2:$H${REL_LAST},\"Luxury\",'All relationships'!$K$2:$K${REL_LAST},\"<>Recently lost\")")
    put(ar, i, 6, f"=COUNTIFS('All relationships'!$A$2:$A${REL_LAST},A{i},'All relationships'!$H$2:$H${REL_LAST},\"Premium\",'All relationships'!$K$2:$K${REL_LAST},\"<>Recently lost\")")
    put(ar, i, 7, lux_txt, align=WRAP)
    put(ar, i, 8, ", ".join(other), align=WRAP)
    put(ar, i, 9, ", ".join(lost), align=WRAP)
    put(ar, i, 10, ", ".join(aka), align=WRAP)
AR_LAST = len(ag_rows) + 1
ar.auto_filter.ref = f"A1:J{AR_LAST}"

# ----- Category by group -----
cg = wb.create_sheet("Category by group", 3)
CATS = list(FIT.keys())
GRP = ["WPP", "Publicis", "Omnicom", "Dentsu", "Havas", "Cheil", "Hakuhodo", "Independent", "Other", "Unknown"]
cg.cell(row=1, column=1, value="Current relationships by category and agency group (excludes Recently lost). Unknown means no agency credited: often in-house, often an open door.").font = BOLD
hdr = [("Category", 28)] + [(g if g != "Unknown" else "Unknown or in-house", 12) for g in GRP] + [("Total", 9)]
header(cg, hdr, row=3)
cg.freeze_panes = "B4"
for i, cat in enumerate(CATS, 4):
    put(cg, i, 1, cat, BOLD)
    for j, g in enumerate(GRP, 2):
        put(cg, i, j, f"=COUNTIFS('All relationships'!$G$2:$G${REL_LAST},$A{i},'All relationships'!$B$2:$B${REL_LAST},\"{g}\",'All relationships'!$K$2:$K${REL_LAST},\"<>Recently lost\")")
    put(cg, i, len(GRP) + 2, f"=SUM(B{i}:{get_column_letter(len(GRP) + 1)}{i})", BOLD)
tot = len(CATS) + 4
put(cg, tot, 1, "Total", BOLD)
for j in range(2, len(GRP) + 3):
    col = get_column_letter(j)
    put(cg, tot, j, f"=SUM({col}4:{col}{tot - 1})", BOLD)

import turmoil
turmoil.add_sheet(wb, recs, header, put, BOLD, WRAP)
import readme
n_a = sum(1 for t in targets if not t["client"] and t["score"] >= 9)
n_b = sum(1 for t in targets if not t["client"] and 7 <= t["score"] < 9)
readme.add_sheet(wb, ASOF, len(recs), len(targets), len(ag_rows), n_a, n_b, FIT, TIER_PTS, BOLD, BODY, WRAP, FONT)

order = ["Read me", "Target list", "In play now", "Agency roster", "Agency turmoil 2025-26", "Category by group", "All relationships"]
wb._sheets = [wb[n] for n in order]
wb.active = 0
for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 100
wb.save(OUT)
print("saved", OUT)
